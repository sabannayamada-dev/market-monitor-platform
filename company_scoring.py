from __future__ import annotations

import argparse
import json
import logging
import math
import os
import random
import re
import string
import sys
import time
import traceback
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from io import BytesIO
from datetime import datetime, timedelta
from html import unescape
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app_meta import APP_BUILD_DATE, APP_NAME, APP_VERSION

try:
    from lxml import etree as LET
except ImportError:
    LET = None


USE_MOCK_DATA = True
USE_JOB_FILE_INPUT = False
USE_HELLOWORK_API = False
USE_EDINET_API = False
USE_EXISTING_CLEAN_DATA = False

JOB_INPUT_FILE = "job_input.xlsx"
EXISTING_CLEAN_DATA_FILE = "all_records_clean.csv"
OUTPUT_DIR = "outputs"
CONFIG_FILE = "scoring_config.json"

REQUEST_INTERVAL_SECONDS = 1.0
REQUEST_TIMEOUT_SECONDS = 30
MAX_RETRIES = 3
ENABLE_API_CACHE = True
CACHE_DIR = ".cache_company_scoring"
EDINET_API_BASE_URL = "https://api.edinet-fsa.go.jp/api/v2"
EDINET_LOOKBACK_DAYS = 400
EDINET_TARGET_DOC_TYPE_CODES = {"120"}

DEFAULT_SCORE_WEIGHTS = {
    "rd": 1.5,
    "growth": 1.2,
    "wage": 1.0,
    "tenure": 1.0,
    "location": 2.0,
}
DEFAULT_TARGET_LOCATIONS = ["金沢", "札幌", "新潟", "関西"]
DEFAULT_MIN_VALID_SCORE_COUNT = 2

MONEY_COLUMNS = [
    "basic_salary_min",
    "basic_salary_max",
    "average_annual_salary",
    "rd_expenses",
    "net_sales",
]
SCORE_COMPONENTS = {
    "rd": ("rd_ratio", "rd_score_norm"),
    "growth": ("salary_growth_potential", "growth_score_norm"),
    "wage": ("real_hourly_wage", "wage_score_norm"),
    "tenure": ("tenure_score", "tenure_score_norm"),
    "location": ("location_match_score", "location_score_norm"),
}

XBRL_TAG_CANDIDATES = {
    "net_sales": [
        "NetSales",
        "NetSalesIFRS",
        "NetSalesSummaryOfBusinessResults",
        "NetSalesIFRSSummaryOfBusinessResults",
        "Revenue",
        "RevenueIFRS",
        "RevenueIFRSSummaryOfBusinessResults",
        "Revenue2IFRS",
        "Revenue2IFRSSummaryOfBusinessResults",
        "RevenueFromContractsWithCustomers",
        "SalesRevenue",
        "OperatingRevenue",
        "OperatingRevenue1",
        "OperatingRevenue2",
        "OperatingRevenueSEC",
        "OperatingRevenueINS",
        "OrdinaryIncome",
        "OrdinaryIncomeBNK",
        "OrdinaryIncomeINS",
        "OrdinaryIncomeLossSummaryOfBusinessResults",
        "OrdinaryIncomeSummaryOfBusinessResults",
    ],
    "rd_expenses": [
        "ResearchAndDevelopmentExpensesResearchAndDevelopmentActivities",
        "ResearchAndDevelopmentExpensesSGA",
        "ResearchAndDevelopmentExpense",
        "ResearchAndDevelopmentExpensesIncludedInGeneralAndAdministrativeExpensesAndManufacturingCostForCurrentPeriod",
        "ResearchAndDevelopmentExpenses",
        "ResearchAndDevelopmentCosts",
    ],
    "average_annual_salary": [
        "AverageAnnualSalaryInformationAboutReportingCompanyInformationAboutEmployees",
        "AverageAnnualSalaryInformationAboutEmployees",
    ],
    "average_length_of_service": [
        "AverageLengthOfServiceYearsInformationAboutReportingCompanyInformationAboutEmployees",
        "AverageLengthOfServiceInformationAboutEmployees",
    ],
    "number_of_employees": [
        "NumberOfEmployees",
        "NumberOfEmployeesInformationAboutReportingCompanyInformationAboutEmployees",
        "NumberOfEmployeesInformationAboutReportingCompanyInformationAboutEmployeesReportingCompany",
        "NumberOfEmployeesInformationAboutEmployees",
        "NumberOfEmployeesSummaryOfBusinessResults",
    ],
    "property_info": [
        "InformationAboutFacilitiesTextBlock",
        "OverviewOfCapitalExpendituresEtcTextBlock",
        "InformationAboutMajorFacilitiesTextBlock",
    ],
    "total_assets": ["Assets", "AssetsSummaryOfBusinessResults", "TotalAssetsSummaryOfBusinessResults", "AssetsIFRS", "TotalAssetsIFRSSummaryOfBusinessResults"],
    "net_assets": [
        "NetAssets",
        "NetAssetsSummaryOfBusinessResults",
        "EquityIFRS",
        "EquityAttributableToOwnersOfParentIFRSSummaryOfBusinessResults",
    ],
    "liabilities": ["Liabilities", "LiabilitiesSummaryOfBusinessResults"],
    "equity_to_asset_ratio": [
        "EquityToAssetRatioSummaryOfBusinessResults",
        "RatioOfOwnersEquityToGrossAssetsIFRSSummaryOfBusinessResults",
    ],
    "operating_income": ["OperatingIncome", "OperatingProfitLossIFRS", "OperatingIncomeSummaryOfBusinessResults"],
    "ordinary_income": ["OrdinaryIncome", "OrdinaryIncomeBNK", "OrdinaryIncomeINS", "OrdinaryIncomeLossSummaryOfBusinessResults"],
    "ordinary_profit": ["OrdinaryProfitLoss", "OrdinaryProfitLossSummaryOfBusinessResults", "OrdinaryIncomeLossSummaryOfBusinessResults"],
    "profit_loss": [
        "ProfitLoss",
        "ProfitLossIFRS",
        "ProfitLossAttributableToOwnersOfParent",
        "ProfitLossAttributableToOwnersOfParentIFRS",
        "ProfitLossAttributableToOwnersOfParentSummaryOfBusinessResults",
        "ProfitLossAttributableToOwnersOfParentIFRSSummaryOfBusinessResults",
    ],
    "gross_profit": ["GrossProfit", "GrossProfitIFRS", "GrossProfitIFRSKeyFinancialData"],
    "cash_flow_operating": [
        "NetCashProvidedByUsedInOperatingActivitiesSummaryOfBusinessResults",
        "CashFlowsFromUsedInOperatingActivitiesIFRS",
        "CashFlowsFromUsedInOperatingActivitiesIFRSSummaryOfBusinessResults",
    ],
    "cash_flow_investing": [
        "NetCashProvidedByUsedInInvestingActivitiesSummaryOfBusinessResults",
        "CashFlowsFromUsedInInvestingActivitiesIFRS",
        "CashFlowsFromUsedInInvestingActivitiesIFRSSummaryOfBusinessResults",
    ],
    "cash_flow_financing": [
        "NetCashProvidedByUsedInFinancingActivitiesSummaryOfBusinessResults",
        "CashFlowsFromUsedInFinancingActivitiesIFRS",
        "CashFlowsFromUsedInFinancingActivitiesIFRSSummaryOfBusinessResults",
    ],
    "cash_and_equivalents": [
        "CashAndCashEquivalentsSummaryOfBusinessResults",
        "CashAndCashEquivalentsIFRS",
        "CashAndCashEquivalentsIFRSSummaryOfBusinessResults",
    ],
    "roe": ["RateOfReturnOnEquitySummaryOfBusinessResults", "RateOfReturnOnEquityIFRSSummaryOfBusinessResults"],
    "eps": [
        "BasicEarningsLossPerShareSummaryOfBusinessResults",
        "BasicEarningsLossPerShareIFRS",
        "BasicEarningsLossPerShareIFRSSummaryOfBusinessResults",
    ],
    "net_assets_per_share": ["NetAssetsPerShareSummaryOfBusinessResults"],
    "dividend_per_share": ["DividendPaidPerShareSummaryOfBusinessResults"],
    "payout_ratio": ["PayoutRatioSummaryOfBusinessResults"],
    "capital_stock": ["CapitalStock", "CapitalStockSummaryOfBusinessResults"],
    "capital_expenditures": ["CapitalExpendituresOverviewOfCapitalExpendituresEtc"],
    "average_age": ["AverageAgeYearsInformationAboutReportingCompanyInformationAboutEmployees"],
    "salary_change_rate": ["ChangeFromPreviousFiscalYearInAverageAnnualSalaryInformationAboutReportingCompanyInformationAboutEmployees"],
    "female_manager_ratio": ["RatioOfFemaleEmployeesInManagerialPositionsMetricsOfReportingCompany"],
    "male_childcare_leave_ratio": [
        "AllEmployeesCalculatedBasedOnProvisionsOfArticle714Item1OfOrdinanceForEnforcementOfActOnChildcareLeaveCaregiverLeaveAndOtherMeasuresForTheWelfareOfWorkersCaringForChildrenOrOtherFamilyMembersRatioOfMaleEmployeesTakingChildcareLeaveMetricsOfReportingCompany"
    ],
    "wage_gap_all": ["AllEmployeesDifferencesInWagesBetweenMaleAndFemaleEmployeesMetricsOfReportingCompany"],
    "wage_gap_regular": ["RegularEmployeesDifferencesInWagesBetweenMaleAndFemaleEmployeesMetricsOfReportingCompany"],
    "wage_gap_non_regular": ["NonRegularEmployeesDifferencesInWagesBetweenMaleAndFemaleEmployeesMetricsOfReportingCompany"],
    "business_description": ["DescriptionOfBusinessTextBlock"],
    "business_risks": ["BusinessRisksTextBlock"],
    "dividend_policy": ["DividendPolicyTextBlock"],
    "rd_activities_text": ["ResearchAndDevelopmentActivitiesTextBlock"],
    "human_resources_policy": [
        "BasicPolicyOnHumanResourcesStrategyEmployeesEtcTextBlock",
        "PolicyOnDevelopmentOfHumanResourcesAndInternalEnvironmentStrategyTextBlock",
    ],
    "product_service_info": ["InformationForEachProductOrServiceTextBlock"],
    "regional_sales_info": ["RevenuesFromExternalCustomersInformationForEachRegionTextBlock"],
}

XBRL_TEXT_FIELDS = {
    "property_info",
    "business_description",
    "business_risks",
    "dividend_policy",
    "rd_activities_text",
    "human_resources_policy",
    "product_service_info",
    "regional_sales_info",
}

XBRL_FIELD_CONTEXT_PREFERENCE = {
    "average_annual_salary": "non_consolidated",
    "average_length_of_service": "non_consolidated",
    "number_of_employees": "non_consolidated",
    "average_age": "non_consolidated",
    "salary_change_rate": "non_consolidated",
    "female_manager_ratio": "non_consolidated",
    "male_childcare_leave_ratio": "non_consolidated",
    "wage_gap_all": "non_consolidated",
    "wage_gap_regular": "non_consolidated",
    "wage_gap_non_regular": "non_consolidated",
}

XBRL_REQUIRED_FIELDS = [
    "net_sales",
    "average_annual_salary",
    "average_length_of_service",
    "rd_expenses",
    "number_of_employees",
    "property_info",
]


class RunContext:
    def __init__(self, run_id: str, output_dir: Path) -> None:
        self.run_id = run_id
        self.output_dir = output_dir
        self.logs_dir = output_dir / "logs"
        self.errors: list[dict[str, Any]] = []
        self.events_path = self.logs_dir / f"{run_id}_events.jsonl"
        self.log_path = self.logs_dir / f"{run_id}.log"


def make_run_id() -> str:
    suffix = "".join(random.choices(string.ascii_lowercase + string.digits, k=6))
    return f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{suffix}"


def ensure_dirs(ctx: RunContext) -> None:
    ctx.output_dir.mkdir(parents=True, exist_ok=True)
    ctx.logs_dir.mkdir(parents=True, exist_ok=True)


def setup_logging(ctx: RunContext) -> None:
    ensure_dirs(ctx)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[
            logging.FileHandler(ctx.log_path, encoding="utf-8-sig"),
            logging.StreamHandler(sys.stdout),
        ],
        force=True,
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="企業評価自動スコアリングシステム")
    parser.add_argument("--version", action="version", version=f"{APP_NAME} v{APP_VERSION} ({APP_BUILD_DATE})")
    parser.add_argument("--rescore", action="store_true", help="既存の all_records_clean.csv を使ってAPIなしで再スコアリングします")
    parser.add_argument("--job-file", default="", help="求人CSV/Excelを読み込んでスコアリングします。EDINETは現時点ではモックデータを使います")
    parser.add_argument("--corporate-master", default="", help="法人番号補完に使うローカルCSV/Excelマスタ")
    parser.add_argument("--use-edinet-api", action="store_true", help="EDINET APIの書類一覧取得を試します。EDINET_API_KEY が必要です")
    parser.add_argument("--edinet-dry-run", action="store_true", help="EDINET書類一覧だけ取得し、XBRL ZIP取得は行いません")
    parser.add_argument("--edinet-lookback-days", type=int, default=EDINET_LOOKBACK_DAYS, help="EDINET書類一覧を遡る日数")
    parser.add_argument("--edinet-xbrl-limit", type=int, default=20, help="EDINET XBRLを取得・解析する最大件数")
    parser.add_argument("--input-clean-data", default=EXISTING_CLEAN_DATA_FILE, help="再スコアリングに使うクレンジング済みCSV")
    parser.add_argument("--config", default=CONFIG_FILE, help="重み設定JSON")
    parser.add_argument("--output-dir", default=OUTPUT_DIR, help="出力フォルダ")
    return parser.parse_args()


def resolve_existing_clean_data_path(path_value: str, output_dir: Path) -> Path:
    path = Path(path_value)
    if path.exists():
        return path
    output_path = output_dir / path_value
    if output_path.exists():
        return output_path
    return path


def to_jsonable(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if pd.isna(value):
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def write_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            clean = {key: to_jsonable(value) for key, value in record.items()}
            handle.write(json.dumps(clean, ensure_ascii=False) + "\n")


def record_event(
    ctx: RunContext,
    level: str,
    event_type: str,
    step: str,
    message: str,
    **kwargs: Any,
) -> None:
    event = {
        "run_id": ctx.run_id,
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "level": level,
        "event_type": event_type,
        "step": step,
        "source": kwargs.get("source"),
        "corporate_number": kwargs.get("corporate_number"),
        "company_name": kwargs.get("company_name"),
        "document_id": kwargs.get("document_id"),
        "message": message,
        "exception_type": kwargs.get("exception_type"),
        "traceback": kwargs.get("traceback"),
    }
    with ctx.events_path.open("a", encoding="utf-8", newline="\n") as handle:
        handle.write(json.dumps(event, ensure_ascii=False) + "\n")
    log_fn = getattr(logging, level.lower(), logging.info)
    log_fn("%s | %s | %s", step, event_type, message)
    if level in {"WARNING", "ERROR"}:
        ctx.errors.append(event)


def cache_path_for(ctx: RunContext, namespace: str, key: str, suffix: str) -> Path:
    safe_key = re.sub(r"[^A-Za-z0-9_.-]+", "_", key)
    path = Path(CACHE_DIR) / namespace / f"{safe_key}.{suffix}"
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def request_json_with_retry(
    url: str,
    params: dict[str, Any],
    ctx: RunContext,
    cache_namespace: str | None = None,
    cache_key: str | None = None,
) -> dict[str, Any] | None:
    cache_path = None
    if ENABLE_API_CACHE and cache_namespace and cache_key:
        cache_path = cache_path_for(ctx, cache_namespace, cache_key, "json")
        if cache_path.exists():
            try:
                record_event(ctx, "INFO", "api_cache_hit", "api", f"キャッシュを使用しました: {cache_path}")
                with cache_path.open("r", encoding="utf-8") as handle:
                    return json.load(handle)
            except Exception as exc:
                record_event(
                    ctx,
                    "WARNING",
                    "api_cache_read_failed",
                    "api",
                    f"キャッシュ読込に失敗したためAPI取得へ進みます: {cache_path}",
                    exception_type=type(exc).__name__,
                    traceback=traceback.format_exc(),
                )

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            record_event(ctx, "INFO", "api_request_started", "api", f"APIリクエスト開始: {url}")
            status_code = 0
            text = ""
            try:
                import requests

                response = requests.get(url, params=params, timeout=REQUEST_TIMEOUT_SECONDS)
                status_code = response.status_code
                text = response.text
            except ImportError:
                request_url = f"{url}?{urllib.parse.urlencode(params)}"
                with urllib.request.urlopen(request_url, timeout=REQUEST_TIMEOUT_SECONDS) as response:
                    status_code = response.status
                    text = response.read().decode("utf-8", errors="replace")
            if status_code == 200:
                data = json.loads(text)
                if cache_path:
                    with cache_path.open("w", encoding="utf-8") as handle:
                        json.dump(data, handle, ensure_ascii=False, indent=2)
                time.sleep(REQUEST_INTERVAL_SECONDS)
                return data
            retryable = status_code in {429, 500, 502, 503, 504}
            record_event(
                ctx,
                "WARNING",
                "api_request_failed",
                "api",
                f"APIステータスが200以外です: {status_code}",
            )
            if not retryable:
                return None
        except urllib.error.HTTPError as exc:
            retryable = exc.code in {429, 500, 502, 503, 504}
            record_event(
                ctx,
                "WARNING",
                "api_request_failed",
                "api",
                f"APIステータスが200以外です: {exc.code}",
                exception_type=type(exc).__name__,
                traceback=traceback.format_exc(),
            )
            if not retryable:
                return None
        except Exception as exc:
            record_event(
                ctx,
                "WARNING",
                "api_request_exception",
                "api",
                "API通信中に例外が発生しました",
                exception_type=type(exc).__name__,
                traceback=traceback.format_exc(),
            )
        time.sleep(REQUEST_INTERVAL_SECONDS * attempt)
    return None


def request_binary_with_retry(
    url: str,
    params: dict[str, Any],
    ctx: RunContext,
    cache_namespace: str | None = None,
    cache_key: str | None = None,
) -> bytes:
    cache_path = None
    if ENABLE_API_CACHE and cache_namespace and cache_key:
        cache_path = cache_path_for(ctx, cache_namespace, cache_key, "bin")
        if cache_path.exists():
            record_event(ctx, "INFO", "api_cache_hit", "api", f"バイナリキャッシュを使用しました: {cache_path}")
            return cache_path.read_bytes()

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            record_event(ctx, "INFO", "api_request_started", "api", f"バイナリAPIリクエスト開始: {url}")
            status_code = 0
            content = b""
            try:
                import requests

                response = requests.get(url, params=params, timeout=REQUEST_TIMEOUT_SECONDS)
                status_code = response.status_code
                content = response.content
            except ImportError:
                request_url = f"{url}?{urllib.parse.urlencode(params)}"
                with urllib.request.urlopen(request_url, timeout=REQUEST_TIMEOUT_SECONDS) as response:
                    status_code = response.status
                    content = response.read()
            if status_code == 200 and content:
                if cache_path:
                    cache_path.write_bytes(content)
                time.sleep(REQUEST_INTERVAL_SECONDS)
                return content
            retryable = status_code in {429, 500, 502, 503, 504}
            record_event(
                ctx,
                "WARNING",
                "api_request_failed",
                "api",
                f"バイナリAPIステータスが200以外です: {status_code}",
            )
            if not retryable:
                return b""
        except urllib.error.HTTPError as exc:
            retryable = exc.code in {429, 500, 502, 503, 504}
            record_event(
                ctx,
                "WARNING",
                "api_request_failed",
                "api",
                f"バイナリAPIステータスが200以外です: {exc.code}",
                exception_type=type(exc).__name__,
                traceback=traceback.format_exc(),
            )
            if not retryable:
                return b""
        except Exception as exc:
            record_event(
                ctx,
                "WARNING",
                "api_request_exception",
                "api",
                "バイナリAPI通信中に例外が発生しました",
                exception_type=type(exc).__name__,
                traceback=traceback.format_exc(),
            )
        time.sleep(REQUEST_INTERVAL_SECONDS * attempt)
    return b""


def edinet_api_key() -> str:
    return os.getenv("EDINET_API_KEY", "")


def load_config(ctx: RunContext, config_file: str = CONFIG_FILE) -> dict[str, Any]:
    config = {
        "score_weights": DEFAULT_SCORE_WEIGHTS.copy(),
        "target_locations": DEFAULT_TARGET_LOCATIONS.copy(),
        "min_valid_score_count": DEFAULT_MIN_VALID_SCORE_COUNT,
    }
    config_path = Path(config_file)
    if not config_path.exists():
        record_event(ctx, "WARNING", "config_missing", "config", "設定ファイルがないためデフォルト設定を使用します")
        return config

    try:
        with config_path.open("r", encoding="utf-8-sig") as handle:
            loaded = json.load(handle)
    except Exception as exc:
        record_event(
            ctx,
            "WARNING",
            "config_load_failed",
            "config",
            "設定ファイルを読めないためデフォルト設定を使用します",
            exception_type=type(exc).__name__,
            traceback=traceback.format_exc(),
        )
        return config

    weights = config["score_weights"].copy()
    loaded_weights = loaded.get("score_weights", {})
    for key, value in loaded_weights.items():
        if key not in weights:
            record_event(ctx, "WARNING", "unknown_weight_key", "config", f"未知の重みキーを無視しました: {key}")
            continue
        try:
            numeric = float(value)
        except (TypeError, ValueError):
            record_event(ctx, "WARNING", "invalid_weight", "config", f"数値ではない重みをデフォルトへ戻しました: {key}")
            continue
        if numeric < 0:
            record_event(ctx, "WARNING", "invalid_weight", "config", f"0未満の重みをデフォルトへ戻しました: {key}")
            continue
        weights[key] = numeric
    if sum(weights.values()) <= 0:
        raise ValueError("全ての重みが0です。スコア計算ができません。")

    config["score_weights"] = weights
    locations = loaded.get("target_locations")
    if isinstance(locations, list) and all(isinstance(item, str) for item in locations):
        config["target_locations"] = locations
    else:
        record_event(ctx, "WARNING", "invalid_locations", "config", "target_locations が不正なためデフォルトを使用します")

    try:
        min_count = int(loaded.get("min_valid_score_count", DEFAULT_MIN_VALID_SCORE_COUNT))
        if min_count < 1:
            raise ValueError
        config["min_valid_score_count"] = min_count
    except (TypeError, ValueError):
        record_event(ctx, "WARNING", "invalid_min_valid_score_count", "config", "min_valid_score_count が不正なためデフォルトを使用します")

    record_event(ctx, "INFO", "config_loaded", "config", "設定ファイルを読み込みました")
    return config


def normalize_corporate_number(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    text = text.translate(str.maketrans("０１２３４５６７８９", "0123456789"))
    text = re.sub(r"[\s\-ー－]", "", text)
    if text.endswith(".0"):
        text = text[:-2]
    return text


def parse_numeric(value: object) -> float:
    if value is None or pd.isna(value):
        return math.nan
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return math.nan
    multiplier = 1.0
    if "百万円" in text:
        multiplier = 1_000_000.0
    elif "千円" in text:
        multiplier = 1_000.0
    text = text.translate(str.maketrans("０１２３４５６７８９．，", "0123456789.,"))
    text = text.replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return math.nan
    return float(match.group(0)) * multiplier


def normalize_text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def load_mock_jobs() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "job_id": "HW-0001",
                "corporate_number": "1010001000001",
                "company_name": "北陸ロボティクス株式会社",
                "office_name": "金沢開発センター",
                "job_title": "機械設計エンジニア",
                "work_location": "石川県金沢市広坂",
                "basic_salary_min": "240,000円",
                "basic_salary_max": "330,000円",
                "annual_holidays": "126日",
                "overtime_hours_avg": "12時間",
                "source": "mock",
            },
            {
                "job_id": "HW-0002",
                "corporate_number": "2020002000002",
                "company_name": "札幌データサイエンス株式会社",
                "office_name": "札幌本社",
                "job_title": "データエンジニア",
                "work_location": "北海道札幌市中央区",
                "basic_salary_min": "260000",
                "basic_salary_max": "390000",
                "annual_holidays": "120",
                "overtime_hours_avg": "18",
                "source": "mock",
            },
            {
                "job_id": "HW-0003",
                "corporate_number": "3030003000003",
                "company_name": "関西バイオマテリアル株式会社",
                "office_name": "大阪研究所",
                "job_title": "研究開発職",
                "work_location": "大阪府吹田市 関西研究拠点",
                "basic_salary_min": "230,000円",
                "basic_salary_max": "310,000円",
                "annual_holidays": "128日",
                "overtime_hours_avg": "8時間",
                "source": "mock",
            },
            {
                "job_id": "HW-0004",
                "corporate_number": "4040004000004",
                "company_name": "新潟精密工業株式会社",
                "office_name": "新潟工場",
                "job_title": "生産技術",
                "work_location": "新潟県長岡市",
                "basic_salary_min": "220000",
                "basic_salary_max": "280000",
                "annual_holidays": "",
                "overtime_hours_avg": "25時間",
                "source": "mock",
            },
            {
                "job_id": "HW-0005",
                "corporate_number": "",
                "company_name": "法人番号未確認テック株式会社",
                "office_name": "東京事業所",
                "job_title": "ソフトウェアエンジニア",
                "work_location": "東京都千代田区",
                "basic_salary_min": "250000",
                "basic_salary_max": "350000",
                "annual_holidays": "122",
                "overtime_hours_avg": "15",
                "source": "mock",
            },
        ]
    )


def load_mock_edinet() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "corporate_number": "1010001000001",
                "edinet_code": "E00001",
                "filer_name": "北陸ロボティクス株式会社",
                "document_id": "S100MOCK1",
                "doc_type_code": "120",
                "fiscal_year_end": "2026-03-31",
                "report_submit_date": "2026-06-28",
                "average_annual_salary": "6,800,000円",
                "average_length_of_service": "9.8年",
                "rd_expenses": "1,200百万円",
                "net_sales": "35,000百万円",
                "number_of_employees": "850人",
                "property_info": "金沢開発センター、富山試験棟、東京営業所",
                "accounting_standard": "JGAAP",
                "edinet_status": "mock",
                "xbrl_parse_status": "mock",
                "xbrl_extracted_fields": "mock",
                "xbrl_missing_fields": "",
                "xbrl_diagnostic_notes": "mock data",
            },
            {
                "corporate_number": "2020002000002",
                "edinet_code": "E00002",
                "filer_name": "札幌データサイエンス株式会社",
                "document_id": "S100MOCK2",
                "doc_type_code": "120",
                "fiscal_year_end": "2026-03-31",
                "report_submit_date": "2026-06-25",
                "average_annual_salary": "7,400,000円",
                "average_length_of_service": "6.5年",
                "rd_expenses": "400百万円",
                "net_sales": "0",
                "number_of_employees": "420人",
                "property_info": "札幌本社、東京開発室",
                "accounting_standard": "JGAAP",
                "edinet_status": "mock",
                "xbrl_parse_status": "mock",
                "xbrl_extracted_fields": "mock",
                "xbrl_missing_fields": "",
                "xbrl_diagnostic_notes": "mock data",
            },
            {
                "corporate_number": "3030003000003",
                "edinet_code": "E00003",
                "filer_name": "関西バイオマテリアル株式会社",
                "document_id": "S100MOCK3",
                "doc_type_code": "120",
                "fiscal_year_end": "2026-03-31",
                "report_submit_date": "2026-06-20",
                "average_annual_salary": "5,900,000円",
                "average_length_of_service": "12.1年",
                "rd_expenses": "2,800百万円",
                "net_sales": "22,000百万円",
                "number_of_employees": "610人",
                "property_info": "関西研究拠点、神戸製造棟",
                "accounting_standard": "JGAAP",
                "edinet_status": "mock",
                "xbrl_parse_status": "mock",
                "xbrl_extracted_fields": "mock",
                "xbrl_missing_fields": "",
                "xbrl_diagnostic_notes": "mock data",
            },
        ]
    )


def fetch_jobs_from_hellowork(ctx: RunContext) -> pd.DataFrame:
    record_event(ctx, "WARNING", "api_stub", "load_jobs", "求人API連携は未実装のスタブです")
    return pd.DataFrame()


def resolve_corporate_number(company_name: str, address: str, ctx: RunContext | None = None) -> dict[str, Any]:
    if ctx:
        record_event(
            ctx,
            "WARNING",
            "api_stub",
            "resolve_corporate_number",
            "法人番号補完APIは未実装のスタブです",
            company_name=company_name,
        )
    return {"corporate_number": "", "status": "not_found", "confidence": 0.0, "notes": "stub"}


def make_recent_date_strings(days: int) -> list[str]:
    today = datetime.now().date()
    return [(today - timedelta(days=offset)).isoformat() for offset in range(days)]


def fetch_edinet_document_list(target_dates: list[str], ctx: RunContext) -> pd.DataFrame:
    api_key = edinet_api_key()
    if not api_key:
        record_event(ctx, "WARNING", "edinet_api_key_missing", "edinet", "EDINET_API_KEY が未設定のためEDINET取得をスキップします")
        return pd.DataFrame()

    records: list[dict[str, Any]] = []
    url = f"{EDINET_API_BASE_URL}/documents.json"
    total_dates = len(target_dates)
    for date_index, date_text in enumerate(target_dates, start=1):
        record_event(ctx, "INFO", "edinet_date_started", "edinet", f"EDINET書類一覧取得中: {date_text} ({date_index}/{total_dates}日目)")
        params = {"date": date_text, "type": 2, "Subscription-Key": api_key}
        data = request_json_with_retry(url, params, ctx, "edinet_document_list", date_text)
        if not data:
            record_event(ctx, "WARNING", "edinet_date_empty", "edinet", f"EDINET書類一覧を取得できませんでした: {date_text} ({date_index}/{total_dates}日目)")
            continue
        before_count = len(records)
        for item in data.get("results", []):
            doc_type_code = normalize_text(item.get("docTypeCode"))
            if doc_type_code not in EDINET_TARGET_DOC_TYPE_CODES:
                continue
            records.append(
                {
                    "corporate_number": normalize_corporate_number(item.get("JCN") or item.get("jcn") or ""),
                    "edinet_code": normalize_text(item.get("edinetCode")),
                    "filer_name": normalize_text(item.get("filerName")),
                    "document_id": normalize_text(item.get("docID")),
                    "doc_type_code": doc_type_code,
                    "sec_code": normalize_text(item.get("secCode")),
                    "fund_code": normalize_text(item.get("fundCode")),
                    "ordinance_code": normalize_text(item.get("ordinanceCode")),
                    "form_code": normalize_text(item.get("formCode")),
                    "period_start": normalize_text(item.get("periodStart")),
                    "fiscal_year_end": normalize_text(item.get("periodEnd")),
                    "report_submit_date": normalize_text(item.get("submitDateTime") or date_text)[:10],
                    "submit_datetime": normalize_text(item.get("submitDateTime")),
                    "doc_description": normalize_text(item.get("docDescription")),
                    "xbrl_flag": normalize_text(item.get("xbrlFlag")),
                    "pdf_flag": normalize_text(item.get("pdfFlag")),
                    "csv_flag": normalize_text(item.get("csvFlag")),
                    "english_doc_flag": normalize_text(item.get("englishDocFlag")),
                    "withdrawal_status": normalize_text(item.get("withdrawalStatus")),
                    "disclosure_status": normalize_text(item.get("disclosureStatus")),
                    "legal_status": normalize_text(item.get("legalStatus")),
                    "average_annual_salary": math.nan,
                    "average_length_of_service": math.nan,
                    "rd_expenses": math.nan,
                    "net_sales": math.nan,
                    "number_of_employees": math.nan,
                    "property_info": "",
                    "accounting_standard": "",
                    "edinet_status": "document_list_only",
                    "xbrl_parse_status": "not_started",
                    "xbrl_extracted_fields": "",
                    "xbrl_missing_fields": "",
                    "xbrl_diagnostic_notes": "",
                }
            )
        added_count = len(records) - before_count
        record_event(ctx, "INFO", "edinet_date_finished", "edinet", f"EDINET書類一覧取得完了: {date_text} ({date_index}/{total_dates}日目) 対象書類 {added_count}件")
    record_event(ctx, "INFO", "edinet_document_list_loaded", "edinet", f"EDINET書類一覧から {len(records)} 件を取得しました")
    return pd.DataFrame(records)


def enrich_edinet_records_with_xbrl(
    df_edinet: pd.DataFrame,
    target_corporate_numbers: set[str],
    ctx: RunContext,
    limit: int,
) -> pd.DataFrame:
    if df_edinet.empty:
        return df_edinet
    df = df_edinet.copy()
    if target_corporate_numbers:
        target_mask = df["corporate_number"].isin(target_corporate_numbers)
        df_targets = df[target_mask].copy()
    else:
        df_targets = df.copy()
    if df_targets.empty:
        record_event(ctx, "WARNING", "edinet_no_target_documents", "edinet", "求人法人番号に一致するEDINET書類候補がありません")
        return df

    df_targets = df_targets.sort_values("report_submit_date", ascending=False).drop_duplicates("corporate_number", keep="first")
    record_event(
        ctx,
        "INFO",
        "edinet_target_documents_selected",
        "edinet",
        f"求人法人番号に一致するEDINET書類候補を {len(df_targets)} 件選定しました",
    )
    if limit > 0:
        df_targets = df_targets.head(limit)

    for idx, row in df_targets.iterrows():
        document_id = normalize_text(row.get("document_id"))
        corp = normalize_corporate_number(row.get("corporate_number"))
        xbrl_bytes = fetch_edinet_xbrl(document_id, ctx)
        values = parse_xbrl_to_edinet_record(xbrl_bytes, corp, ctx)
        if not values:
            df.loc[idx, "edinet_status"] = "xbrl_not_parsed"
            df.loc[idx, "xbrl_parse_status"] = "not_parsed"
            df.loc[idx, "xbrl_diagnostic_notes"] = "XBRL ZIPを取得できない、または解析できませんでした"
            continue
        for field, value in values.items():
            if field == "_xbrl_extracted_fields":
                df.loc[idx, "xbrl_extracted_fields"] = value
                continue
            if field == "_xbrl_missing_fields":
                df.loc[idx, "xbrl_missing_fields"] = value
                continue
            df.loc[idx, field] = value
        df.loc[idx, "edinet_status"] = "xbrl_parsed"
        df.loc[idx, "xbrl_parse_status"] = "parsed"
        missing = normalize_text(df.loc[idx].get("xbrl_missing_fields", ""))
        if missing:
            df.loc[idx, "xbrl_diagnostic_notes"] = f"一部XBRLタグを抽出できませんでした: {missing}"
        else:
            df.loc[idx, "xbrl_diagnostic_notes"] = "XBRLタグ抽出が完了しました"
    return df


def export_edinet_snapshot(df_edinet: pd.DataFrame, ctx: RunContext) -> None:
    if df_edinet.empty:
        return
    snapshot_columns = [
        "corporate_number",
        "edinet_code",
        "filer_name",
        "document_id",
        "doc_type_code",
        "fiscal_year_end",
        "report_submit_date",
        "edinet_status",
        "xbrl_parse_status",
        "xbrl_extracted_fields",
        "xbrl_missing_fields",
    ]
    columns = [column for column in snapshot_columns if column in df_edinet.columns]
    path = ctx.output_dir / "edinet_documents_snapshot.csv"
    df_edinet[columns].to_csv(path, index=False, encoding="utf-8-sig")
    record_event(ctx, "INFO", "edinet_snapshot_exported", "edinet", f"EDINET書類候補スナップショットを出力しました: {path}")


def fetch_edinet_xbrl(document_id: str, ctx: RunContext) -> bytes:
    api_key = edinet_api_key()
    if not api_key:
        record_event(ctx, "WARNING", "edinet_api_key_missing", "edinet", "EDINET_API_KEY が未設定のためXBRL取得をスキップします", document_id=document_id)
        return b""
    if not document_id:
        return b""
    url = f"{EDINET_API_BASE_URL}/documents/{document_id}"
    params = {"type": 1, "Subscription-Key": api_key}
    return request_binary_with_retry(url, params, ctx, "edinet_xbrl", document_id)


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag.split(":")[-1]


def clean_text_block(value: str) -> str:
    text = unescape(value or "")
    text = re.sub(r"<[^>]+>", " ", text)
    return normalize_text(text)


def context_priority(context_ref: str | None, field: str = "") -> int:
    text = (context_ref or "").lower()
    score = 0
    if "current" in text or "currentyear" in text:
        score += 50
    if "duration" in text:
        score += 30
    if "instant" in text:
        score += 10
    is_non_consolidated = "nonconsolidated" in text or "non_consolidated" in text
    is_consolidated = ("consolidated" in text or "consolidatedmember" in text) and not is_non_consolidated
    has_dimension_member = "_" in text or "member" in text
    is_plain_period_context = not has_dimension_member
    if XBRL_FIELD_CONTEXT_PREFERENCE.get(field) == "non_consolidated":
        if is_non_consolidated:
            score += 16
        elif is_consolidated:
            score -= 4
    else:
        if is_plain_period_context:
            score += 20
        if is_consolidated:
            score += 10
        if is_non_consolidated:
            score -= 12
        if has_dimension_member:
            score -= 10
        if any(token in text for token in ["reportablesegment", "operatingsegments", "reconcilingitems"]):
            score -= 12
        if "totalofreportablesegments" in text:
            score += 4
    if "filingdate" in text:
        score -= 20
    return score


def extract_employee_table_values(html_text: str) -> dict[str, Any]:
    cells: list[str] = []
    for match in re.finditer(r"<t[dh][^>]*>(.*?)</t[dh]>", html_text or "", flags=re.IGNORECASE | re.DOTALL):
        cell = re.sub(r"<[^>]+>", " ", match.group(1))
        cell = normalize_text(cell)
        if cell:
            cells.append(cell)
    if not cells:
        return {}

    values: dict[str, Any] = {}
    position_candidates: dict[str, list[int]] = {
        "number_of_employees": [],
        "average_length_of_service": [],
        "average_annual_salary": [],
    }
    for index, cell in enumerate(cells):
        if "従業員数" in cell and "人" in cell and "average" not in cell.lower():
            position_candidates["number_of_employees"].append(index)
        if "平均勤続年数" in cell:
            position_candidates["average_length_of_service"].append(index)
        if "平均年間給与" in cell and "増減率" not in cell:
            position_candidates["average_annual_salary"].append(index)

    if not any(position_candidates.values()):
        return {}
    header_positions: dict[str, int] = {}
    salary_positions = position_candidates["average_annual_salary"]
    if salary_positions:
        salary_position = salary_positions[0]
        header_positions["average_annual_salary"] = salary_position
        for field in ["number_of_employees", "average_length_of_service"]:
            before_salary = [position for position in position_candidates[field] if position < salary_position and salary_position - position <= 10]
            if before_salary:
                header_positions[field] = max(before_salary)
    else:
        for field, positions in position_candidates.items():
            if positions:
                header_positions[field] = positions[0]

    first_header = min(header_positions.values())
    row_start = -1
    header_words = ["従業員", "平均", "給与", "年齢", "勤続", "増減率", "令和", "現在"]
    for index in range(first_header + 1, len(cells)):
        if any(word in cells[index] for word in header_words):
            continue
        parsed = parse_numeric(cells[index])
        if not pd.isna(parsed):
            row_start = index
            break
    if row_start < 0:
        return {}
    data_cells = [cell for cell in cells[row_start:] if not re.fullmatch(r"[\[〔（(].*[\]〕）)]", cell)]
    for field, header_index in header_positions.items():
        value_index = header_index - first_header
        if value_index < len(data_cells):
            parsed = parse_numeric(data_cells[value_index])
            if not pd.isna(parsed):
                header_cell = cells[header_index]
                if field == "average_annual_salary":
                    if "千円" in header_cell:
                        parsed *= 1000
                    if parsed < 100_000:
                        continue
                if field == "average_length_of_service" and parsed > 60:
                    continue
                values[field] = parsed
    return values


def parse_xml_root(xml_bytes: bytes) -> Any | None:
    if LET is not None:
        try:
            return LET.fromstring(xml_bytes, parser=LET.XMLParser(recover=True, huge_tree=True))
        except LET.XMLSyntaxError:
            return None
    try:
        return ET.fromstring(xml_bytes)
    except ET.ParseError:
        return None


def extract_xbrl_values_from_xml(xml_bytes: bytes) -> dict[str, Any]:
    found: dict[str, tuple[int, str]] = {}
    tag_to_field: dict[str, str] = {}
    for field, candidates in XBRL_TAG_CANDIDATES.items():
        for candidate in candidates:
            tag_to_field[candidate] = field

    root = parse_xml_root(xml_bytes)
    if root is None:
        return {}

    for element in root.iter():
        name = local_name(str(element.tag))
        context_ref = element.attrib.get("contextRef")
        field = tag_to_field.get(name)
        if not field and name != "InformationAboutEmployeesTextBlock":
            continue

        raw_text = normalize_text("".join(element.itertext()))
        if not raw_text:
            continue
        if name == "InformationAboutEmployeesTextBlock":
            for fallback_field, fallback_value in extract_employee_table_values(raw_text).items():
                previous = found.get(fallback_field)
                fallback_priority = context_priority(context_ref, fallback_field)
                if previous is None or fallback_priority >= previous[0]:
                    found[fallback_field] = (fallback_priority, str(fallback_value))
        if not field:
            continue
        priority = context_priority(context_ref, field)
        previous = found.get(field)
        if previous is None or priority > previous[0]:
            found[field] = (priority, raw_text)

    values: dict[str, Any] = {}
    for field, (_, raw_text) in found.items():
        if field in XBRL_TEXT_FIELDS:
            values[field] = clean_text_block(raw_text)
        else:
            values[field] = parse_numeric(raw_text)
    return values


def parse_xbrl_to_edinet_record(xbrl_zip_bytes: bytes, corporate_number: str, ctx: RunContext) -> dict[str, Any]:
    if not xbrl_zip_bytes:
        return {}
    values: dict[str, Any] = {}
    try:
        with zipfile.ZipFile(BytesIO(xbrl_zip_bytes)) as archive:
            names = [
                name
                for name in archive.namelist()
                if name.lower().endswith((".xbrl", ".xml")) and ("xbrl" in name.lower() or "publicdoc" in name.lower())
            ]
            for name in names:
                try:
                    extracted = extract_xbrl_values_from_xml(archive.read(name))
                except Exception as exc:
                    record_event(
                        ctx,
                        "WARNING",
                        "xbrl_parse_failed",
                        "edinet",
                        f"XBRLファイルの解析に失敗しました: {name}",
                        corporate_number=corporate_number,
                        exception_type=type(exc).__name__,
                        traceback=traceback.format_exc(),
                    )
                    continue
                for field, value in extracted.items():
                    if field not in values or pd.isna(values.get(field)) or not values.get(field):
                        values[field] = value
                if all(field in values for field in XBRL_REQUIRED_FIELDS):
                    break
    except zipfile.BadZipFile as exc:
        record_event(
            ctx,
            "WARNING",
            "xbrl_zip_invalid",
            "edinet",
            "EDINETから取得したXBRL ZIPを開けませんでした",
            corporate_number=corporate_number,
            exception_type=type(exc).__name__,
            traceback=traceback.format_exc(),
        )
        return {}

    missing = [field for field in XBRL_REQUIRED_FIELDS if field not in values]
    values["_xbrl_extracted_fields"] = ",".join(sorted(field for field in values if not field.startswith("_")))
    values["_xbrl_missing_fields"] = ",".join(missing)
    if missing:
        record_event(
            ctx,
            "WARNING",
            "xbrl_fields_missing",
            "edinet",
            f"XBRLから取得できない項目があります: {', '.join(missing)}",
            corporate_number=corporate_number,
        )
    else:
        record_event(ctx, "INFO", "xbrl_parsed", "edinet", "XBRL解析が完了しました", corporate_number=corporate_number)
    return values


def load_existing_clean_data(path: Path, ctx: RunContext) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"再スコアリング用ファイルが見つかりません: {path}")
    df = pd.read_csv(path, encoding="utf-8-sig", dtype={"corporate_number": "string"})
    df["corporate_number"] = df["corporate_number"].fillna("").map(normalize_corporate_number)

    numeric_columns = [
        "basic_salary_min",
        "basic_salary_max",
        "annual_holidays",
        "overtime_hours_avg",
        "average_annual_salary",
        "average_length_of_service",
        "rd_expenses",
        "net_sales",
        "number_of_employees",
    ]
    for column in numeric_columns:
        if column in df.columns:
            df[column] = df[column].map(parse_numeric)

    text_columns = [
        "job_id",
        "company_name",
        "office_name",
        "job_title",
        "work_location",
        "source",
        "filer_name",
        "property_info",
        "accounting_standard",
        "join_status",
        "score_status",
        "missing_fields",
        "data_quality_notes",
    ]
    for column in text_columns:
        if column in df.columns:
            df[column] = df[column].map(normalize_text)

    required_columns = ["join_status", "data_quality_notes"]
    for column in required_columns:
        if column not in df.columns:
            df[column] = ""
            record_event(ctx, "WARNING", "field_missing", "load_existing_clean_data", f"再スコアリング用データに列がないため補完しました: {column}")

    record_event(ctx, "INFO", "existing_clean_data_loaded", "load", f"再スコアリング用データを読み込みました: {path}")
    return df


JOB_COLUMN_ALIASES = {
    "job_id": ["job_id", "求人番号", "求人No", "求人NO", "求人id", "求人ID", "kyujin_bangho"],
    "corporate_number": ["corporate_number", "法人番号", "法人番号13桁", "hojin_bangho"],
    "company_name": ["company_name", "企業名", "会社名", "事業所名", "jigyosho_mei"],
    "office_name": ["office_name", "事業所名", "拠点名", "office"],
    "job_title": ["job_title", "職種", "募集職種", "shokushu"],
    "work_location": ["work_location", "就業場所", "勤務地", "就業場所住所", "shugyo_basho_jusho"],
    "basic_salary_min": ["basic_salary_min", "基本給下限", "基本給_下限", "月給下限", "kihonkyu_kaishi"],
    "basic_salary_max": ["basic_salary_max", "基本給上限", "基本給_上限", "月給上限", "kihonkyu_shuryo"],
    "annual_holidays": ["annual_holidays", "年間休日数", "年間休日", "nenkan_kyujitsu"],
    "overtime_hours_avg": ["overtime_hours_avg", "月平均残業時間", "平均残業時間", "月平均時間外労働", "jikangai_rodo_avg"],
}

COLLECTION_METADATA_COLUMNS = [
    "source_id",
    "source_record_id",
    "source_quality_status",
    "source_quality_notes",
]

CORPORATE_MASTER_COLUMN_ALIASES = {
    "corporate_number": ["corporate_number", "法人番号", "法人番号13桁", "hojin_bangho"],
    "company_name": ["company_name", "企業名", "会社名", "商号又は名称", "name"],
    "address": ["address", "所在地", "住所", "本店又は主たる事務所の所在地", "location"],
}


def normalize_column_label(label: object) -> str:
    text = normalize_text(label)
    text = text.translate(str.maketrans("０１２３４５６７８９ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺａｂｃｄｅｆｇｈｉｊｋｌｍｎｏｐｑｒｓｔｕｖｗｘｙｚ", "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz"))
    return re.sub(r"[\s_\-・　（）()\[\]【】]", "", text).lower()


def build_job_column_map(columns: list[str], ctx: RunContext) -> dict[str, str]:
    normalized_columns = {normalize_column_label(column): column for column in columns}
    mapping: dict[str, str] = {}
    used_source_columns: set[str] = set()
    for target, aliases in JOB_COLUMN_ALIASES.items():
        for alias in aliases:
            source = normalized_columns.get(normalize_column_label(alias))
            if source and source not in used_source_columns:
                mapping[source] = target
                used_source_columns.add(source)
                break
    ignored_columns = set(COLLECTION_METADATA_COLUMNS)
    unmapped = [column for column in columns if column not in used_source_columns and column not in ignored_columns]
    if unmapped:
        record_event(ctx, "WARNING", "unmapped_input_columns", "load_jobs_from_file", f"未使用の入力列があります: {', '.join(map(str, unmapped))}")
    missing_required = [column for column in ["job_id", "company_name"] if column not in mapping.values()]
    if missing_required:
        record_event(ctx, "WARNING", "required_input_columns_missing", "load_jobs_from_file", f"推奨列が不足しています: {', '.join(missing_required)}")
    return mapping


def build_alias_column_map(columns: list[str], aliases: dict[str, list[str]], ctx: RunContext, step: str) -> dict[str, str]:
    normalized_columns = {normalize_column_label(column): column for column in columns}
    mapping: dict[str, str] = {}
    used_source_columns: set[str] = set()
    for target, candidates in aliases.items():
        for alias in candidates:
            source = normalized_columns.get(normalize_column_label(alias))
            if source and source not in used_source_columns:
                mapping[source] = target
                used_source_columns.add(source)
                break
    unmapped = [column for column in columns if column not in used_source_columns]
    if unmapped:
        record_event(ctx, "WARNING", "unmapped_input_columns", step, f"未使用の入力列があります: {', '.join(map(str, unmapped))}")
    return mapping


def load_jobs_from_file(path: Path, ctx: RunContext) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"求人入力ファイルが見つかりません: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        df_raw = pd.read_csv(path, encoding="utf-8-sig", dtype=str)
    elif suffix in {".xlsx", ".xlsm", ".xls"}:
        df_raw = pd.read_excel(path, dtype=str)
    else:
        raise ValueError(f"未対応の求人入力ファイル形式です: {path.suffix}")

    column_map = build_job_column_map(list(df_raw.columns), ctx)
    df = df_raw.rename(columns=column_map)
    keep_columns = list(JOB_COLUMN_ALIASES.keys())
    for column in COLLECTION_METADATA_COLUMNS:
        if column in df_raw.columns and column not in keep_columns:
            keep_columns.append(column)
    for column in keep_columns:
        if column not in df.columns:
            df[column] = ""
    df = df[keep_columns].copy()
    df["source"] = "file"
    record_event(ctx, "INFO", "input_loaded", "load_jobs_from_file", f"求人ファイルを読み込みました: {path}")
    return df


def load_corporate_master(path: Path, ctx: RunContext) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"法人番号マスタが見つかりません: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        df_raw = pd.read_csv(path, encoding="utf-8-sig", dtype=str)
    elif suffix in {".xlsx", ".xlsm", ".xls"}:
        df_raw = pd.read_excel(path, dtype=str)
    else:
        raise ValueError(f"未対応の法人番号マスタ形式です: {path.suffix}")

    column_map = build_alias_column_map(list(df_raw.columns), CORPORATE_MASTER_COLUMN_ALIASES, ctx, "load_corporate_master")
    df = df_raw.rename(columns=column_map)
    for column in CORPORATE_MASTER_COLUMN_ALIASES:
        if column not in df.columns:
            df[column] = ""
    df = df[list(CORPORATE_MASTER_COLUMN_ALIASES.keys())].copy()
    df["corporate_number"] = df["corporate_number"].map(normalize_corporate_number)
    df["company_name"] = df["company_name"].map(normalize_text)
    df["address"] = df["address"].map(normalize_text)
    df = df[df["corporate_number"].map(lambda value: bool(re.fullmatch(r"\d{13}", value or "")))].copy()
    record_event(ctx, "INFO", "corporate_master_loaded", "corporate_number", f"法人番号マスタを読み込みました: {path} ({len(df)}件)")
    return df


def normalize_match_text(value: object) -> str:
    text = normalize_text(value)
    text = text.translate(str.maketrans("０１２３４５６７８９ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺａｂｃｄｅｆｇｈｉｊｋｌｍｎｏｐｑｒｓｔｕｖｗｘｙｚ", "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz"))
    for token in ["株式会社", "有限会社", "合同会社", "（株）", "(株)", "㈱", "　", " ", "-", "－", "ー"]:
        text = text.replace(token, "")
    return text.lower()


def address_overlap_score(job_location: str, master_address: str) -> float:
    job = normalize_match_text(job_location)
    master = normalize_match_text(master_address)
    if not job or not master:
        return 0.0
    if job in master or master in job:
        return 1.0
    common = 0
    for token_len in [6, 5, 4, 3]:
        for start in range(max(len(job) - token_len + 1, 0)):
            if job[start : start + token_len] in master:
                common += token_len
                break
        if common:
            break
    return min(common / max(len(job), 1), 1.0)


def find_corporate_number_candidates(row: pd.Series, master: pd.DataFrame) -> list[dict[str, Any]]:
    company_name = normalize_text(row.get("company_name", ""))
    work_location = normalize_text(row.get("work_location", ""))
    name_key = normalize_match_text(company_name)
    if not name_key:
        return []

    candidates: list[dict[str, Any]] = []
    for _, item in master.iterrows():
        master_name = normalize_text(item.get("company_name", ""))
        master_name_key = normalize_match_text(master_name)
        if not master_name_key:
            continue
        if name_key == master_name_key:
            name_score = 1.0
        elif name_key in master_name_key or master_name_key in name_key:
            name_score = 0.82
        else:
            continue
        address_score = address_overlap_score(work_location, normalize_text(item.get("address", "")))
        confidence = round((name_score * 0.75) + (address_score * 0.25), 3)
        candidates.append(
            {
                "corporate_number": item.get("corporate_number", ""),
                "company_name": master_name,
                "address": normalize_text(item.get("address", "")),
                "confidence": confidence,
                "name_score": name_score,
                "address_score": round(address_score, 3),
            }
        )
    return sorted(candidates, key=lambda item: item["confidence"], reverse=True)


def apply_corporate_number_master(df_jobs: pd.DataFrame, master: pd.DataFrame, ctx: RunContext) -> pd.DataFrame:
    df = df_jobs.copy()
    for column in ["corporate_number_resolution_status", "corporate_number_resolution_confidence", "corporate_number_candidates", "corporate_number_resolution_notes"]:
        if column not in df.columns:
            df[column] = ""
    df["corporate_number_resolution_confidence"] = pd.to_numeric(df["corporate_number_resolution_confidence"], errors="coerce")

    for idx, row in df.iterrows():
        current = normalize_corporate_number(row.get("corporate_number"))
        if re.fullmatch(r"\d{13}", current or ""):
            df.loc[idx, "corporate_number_resolution_status"] = "already_present"
            df.loc[idx, "corporate_number_resolution_confidence"] = 1.0
            continue
        candidates = find_corporate_number_candidates(row, master)
        if not candidates:
            df.loc[idx, "corporate_number_resolution_status"] = "not_found"
            df.loc[idx, "corporate_number_resolution_notes"] = "法人番号マスタに候補がありません"
            continue
        df.loc[idx, "corporate_number_candidates"] = json.dumps(candidates[:5], ensure_ascii=False)
        top = candidates[0]
        second_confidence = candidates[1]["confidence"] if len(candidates) > 1 else 0.0
        if top["confidence"] >= 0.92 and (top["confidence"] - second_confidence) >= 0.05:
            df.loc[idx, "corporate_number"] = top["corporate_number"]
            df.loc[idx, "corporate_number_resolution_status"] = "estimated"
            df.loc[idx, "corporate_number_resolution_confidence"] = top["confidence"]
            df.loc[idx, "corporate_number_resolution_notes"] = f"法人番号マスタから高信頼で補完: {top['company_name']}"
            record_event(ctx, "INFO", "corporate_number_resolved", "corporate_number", f"法人番号を補完しました: {top['corporate_number']}", company_name=row.get("company_name"))
        else:
            df.loc[idx, "corporate_number_resolution_status"] = "ambiguous"
            df.loc[idx, "corporate_number_resolution_confidence"] = top["confidence"]
            df.loc[idx, "corporate_number_resolution_notes"] = "候補が曖昧なため自動採用しません"
            record_event(ctx, "WARNING", "corporate_number_ambiguous", "corporate_number", "法人番号候補が曖昧です", company_name=row.get("company_name"))
    return df


def clean_jobs(df: pd.DataFrame, ctx: RunContext) -> pd.DataFrame:
    df = df.copy()
    expected = [
        "job_id",
        "corporate_number",
        "company_name",
        "office_name",
        "job_title",
        "work_location",
        "basic_salary_min",
        "basic_salary_max",
        "annual_holidays",
        "overtime_hours_avg",
        "source",
        *COLLECTION_METADATA_COLUMNS,
    ]
    for column in expected:
        if column not in df.columns:
            df[column] = ""
            if column not in COLLECTION_METADATA_COLUMNS:
                record_event(ctx, "WARNING", "field_missing", "clean_jobs", f"求人データに列がないため空欄を補完しました: {column}")
    resolution_columns = [
        "corporate_number_resolution_status",
        "corporate_number_resolution_confidence",
        "corporate_number_candidates",
        "corporate_number_resolution_notes",
    ]
    for column in resolution_columns:
        if column not in df.columns:
            df[column] = ""
    df["corporate_number"] = df["corporate_number"].map(normalize_corporate_number)
    df["corporate_number_status"] = df["corporate_number"].map(
        lambda value: "ok" if re.fullmatch(r"\d{13}", value or "") else ("missing" if not value else "invalid")
    )
    df.loc[df["corporate_number_resolution_status"].eq("estimated"), "corporate_number_status"] = "estimated"
    for column in ["basic_salary_min", "basic_salary_max", "annual_holidays", "overtime_hours_avg"]:
        df[column] = df[column].map(parse_numeric)
    for column in ["job_id", "company_name", "office_name", "job_title", "work_location", "source", *COLLECTION_METADATA_COLUMNS]:
        df[column] = df[column].map(normalize_text)
    df["data_quality_notes"] = ""
    return df[expected + ["corporate_number_status", *resolution_columns, "data_quality_notes"]]


def clean_edinet(df: pd.DataFrame, ctx: RunContext) -> pd.DataFrame:
    df = df.copy()
    expected = [
        "corporate_number",
        "edinet_code",
        "filer_name",
        "document_id",
        "doc_type_code",
        "fiscal_year_end",
        "report_submit_date",
        "average_annual_salary",
        "average_length_of_service",
        "rd_expenses",
        "net_sales",
        "number_of_employees",
        "property_info",
        "accounting_standard",
        "edinet_status",
        "xbrl_parse_status",
        "xbrl_extracted_fields",
        "xbrl_missing_fields",
        "xbrl_diagnostic_notes",
    ]
    for column in expected:
        if column not in df.columns:
            df[column] = ""
            record_event(ctx, "WARNING", "field_missing", "clean_edinet", f"EDINETデータに列がないため空欄を補完しました: {column}")
    df["corporate_number"] = df["corporate_number"].map(normalize_corporate_number)
    for column in [
        "average_annual_salary",
        "average_length_of_service",
        "rd_expenses",
        "net_sales",
        "number_of_employees",
    ]:
        df[column] = df[column].map(parse_numeric)
    for column in [
        "edinet_code",
        "filer_name",
        "document_id",
        "doc_type_code",
        "fiscal_year_end",
        "report_submit_date",
        "property_info",
        "accounting_standard",
        "edinet_status",
        "xbrl_parse_status",
        "xbrl_extracted_fields",
        "xbrl_missing_fields",
        "xbrl_diagnostic_notes",
    ]:
        df[column] = df[column].map(normalize_text)
    return df[expected]


def merge_jobs_and_edinet(df_jobs: pd.DataFrame, df_edinet: pd.DataFrame) -> pd.DataFrame:
    duplicate_keys = set(df_edinet[df_edinet.duplicated("corporate_number", keep=False)]["corporate_number"].dropna())
    df = df_jobs.merge(df_edinet, on="corporate_number", how="left", suffixes=("_job", "_edinet"))

    def status(row: pd.Series) -> str:
        corp = row.get("corporate_number", "")
        if not corp:
            return "no_corporate_number"
        if corp in duplicate_keys:
            return "duplicate_match"
        if pd.isna(row.get("edinet_code")) or not str(row.get("edinet_code", "")).strip():
            return "no_edinet_match"
        return "matched"

    df["join_status"] = df.apply(status, axis=1)

    def failure_reason(row: pd.Series) -> str:
        join_status = row.get("join_status")
        corp = row.get("corporate_number", "")
        corp_status = row.get("corporate_number_status", "")
        if join_status == "matched":
            return "not_failed"
        if join_status == "no_corporate_number":
            return "missing_corporate_number"
        if corp_status == "invalid" or (corp and not re.fullmatch(r"\d{13}", str(corp))):
            return "invalid_corporate_number_format"
        if join_status == "duplicate_match":
            return "edinet_duplicate_candidates"
        if join_status == "no_edinet_match":
            return "edinet_record_not_found"
        return "unknown"

    def diagnostic_notes(row: pd.Series) -> str:
        reason = row.get("join_failure_reason")
        corp = row.get("corporate_number", "")
        company = row.get("company_name", "")
        if reason == "not_failed":
            return ""
        if reason == "missing_corporate_number":
            return f"求人データに法人番号がありません。企業名・所在地による法人番号補完候補です: {company}"
        if reason == "invalid_corporate_number_format":
            return f"法人番号が13桁数字ではありません: {corp}"
        if reason == "edinet_duplicate_candidates":
            return f"EDINET側に同じ法人番号の候補が複数あります: {corp}"
        if reason == "edinet_record_not_found":
            return f"EDINETデータに法人番号一致がありません。未上場、EDINET非提出、法人番号違いの可能性があります: {corp}"
        return "JOIN失敗理由を特定できませんでした"

    df["join_failure_reason"] = df.apply(failure_reason, axis=1)
    df["join_diagnostic_notes"] = df.apply(diagnostic_notes, axis=1)
    df["edinet_status"] = df["edinet_status"].fillna("not_found")
    for column in ["filer_name", "property_info", "accounting_standard", "edinet_code", "document_id", "doc_type_code"]:
        if column in df.columns:
            df[column] = df[column].fillna("")

    def name_check_status(row: pd.Series) -> str:
        if row.get("join_status") != "matched":
            return "not_checked"
        company_key = normalize_match_text(row.get("company_name", ""))
        filer_key = normalize_match_text(row.get("filer_name", ""))
        if not company_key or not filer_key:
            return "not_checked"
        if company_key == filer_key or company_key in filer_key or filer_key in company_key:
            return "ok"
        if row.get("corporate_number_status") == "estimated":
            return "estimated_number_name_mismatch"
        return "name_mismatch"

    def name_check_notes(row: pd.Series) -> str:
        status = row.get("join_name_check_status", "")
        if status in {"ok", "not_checked"}:
            return ""
        return f"求人企業名とEDINET提出者名が一致しません。求人='{row.get('company_name', '')}', EDINET='{row.get('filer_name', '')}'"

    df["join_name_check_status"] = df.apply(name_check_status, axis=1)
    df["join_name_check_notes"] = df.apply(name_check_notes, axis=1)
    return df


def safe_divide(numerator: Any, denominator: Any) -> float:
    if pd.isna(numerator) or pd.isna(denominator):
        return math.nan
    if denominator <= 0:
        return math.nan
    return float(numerator) / float(denominator)


def count_location_matches(row: pd.Series, locations: list[str]) -> int:
    haystack = normalize_text(f"{row.get('work_location', '')} {row.get('property_info', '')}")
    normalized = haystack.replace(" ", "")
    count = 0
    for keyword in locations:
        if keyword.replace(" ", "") in normalized:
            count += 1
    return count


def normalize_score(series: pd.Series) -> pd.Series:
    numeric = pd.to_numeric(series, errors="coerce")
    valid = numeric.dropna()
    if len(valid) < 2:
        return numeric.map(lambda value: 50.0 if pd.notna(value) else math.nan)
    lower = valid.quantile(0.05)
    upper = valid.quantile(0.95)
    if upper == lower:
        return numeric.map(lambda value: 50.0 if pd.notna(value) else math.nan)
    return ((numeric - lower) / (upper - lower) * 100).clip(lower=0, upper=100)


def calculate_score_columns(df: pd.DataFrame, config: dict[str, Any], ctx: RunContext) -> pd.DataFrame:
    df = df.copy()
    net_sales = pd.to_numeric(df["net_sales"], errors="coerce")
    rd_expenses = pd.to_numeric(df["rd_expenses"], errors="coerce")
    average_salary = pd.to_numeric(df["average_annual_salary"], errors="coerce")
    basic_salary_min = pd.to_numeric(df["basic_salary_min"], errors="coerce")
    annual_holidays = pd.to_numeric(df["annual_holidays"], errors="coerce")
    overtime_hours = pd.to_numeric(df["overtime_hours_avg"], errors="coerce")
    average_tenure = pd.to_numeric(df["average_length_of_service"], errors="coerce")

    df["rd_ratio"] = (rd_expenses / net_sales.where(net_sales > 0)).replace([math.inf, -math.inf], math.nan)
    annual_min_salary = basic_salary_min * 12
    df["salary_growth_potential"] = (average_salary / annual_min_salary.where(annual_min_salary > 0)).replace([math.inf, -math.inf], math.nan)
    valid_hours = annual_holidays.notna() & overtime_hours.notna() & annual_holidays.ge(0) & annual_holidays.lt(365)
    annual_work_hours = (365 - annual_holidays) * 8 + overtime_hours * 12
    df["annual_work_hours"] = annual_work_hours.where(valid_hours)
    df["real_hourly_wage"] = (average_salary / df["annual_work_hours"].where(df["annual_work_hours"] > 0)).replace([math.inf, -math.inf], math.nan)
    df["tenure_score"] = average_tenure.where(average_tenure >= 0)
    df["location_match_score"] = df.apply(lambda row: count_location_matches(row, config["target_locations"]), axis=1)

    ranking_candidates = df["join_status"].eq("matched")
    for raw_column, norm_column in SCORE_COMPONENTS.values():
        norm = pd.Series(math.nan, index=df.index, dtype="float64")
        norm.loc[ranking_candidates] = normalize_score(df.loc[ranking_candidates, raw_column])
        df[norm_column] = norm

    weights = config["score_weights"]
    total_scores: list[float] = []
    score_statuses: list[str] = []
    missing_fields_values: list[str] = []
    notes_values: list[str] = []
    for _, row in df.iterrows():
        weighted_sum = 0.0
        weight_sum = 0.0
        missing: list[str] = []
        valid_count = 0
        for key, (_, norm_column) in SCORE_COMPONENTS.items():
            value = row.get(norm_column)
            if pd.isna(value):
                missing.append(norm_column)
                continue
            weight = weights[key]
            weighted_sum += float(value) * weight
            weight_sum += weight
            valid_count += 1
        if row.get("join_status") != "matched":
            total_scores.append(math.nan)
            score_statuses.append("insufficient_data")
        elif valid_count < config["min_valid_score_count"] or weight_sum <= 0:
            total_scores.append(math.nan)
            score_statuses.append("insufficient_data")
        else:
            total_scores.append(weighted_sum / weight_sum)
            score_statuses.append("ok" if not missing else "partial")
        missing_fields_values.append(",".join(missing))
        notes = normalize_text(row.get("data_quality_notes", ""))
        if row.get("join_name_check_status") == "estimated_number_name_mismatch":
            notes = (notes + "; " if notes else "") + "estimated corporate number joined to EDINET but company name differs"
        if pd.notna(row.get("salary_growth_potential")):
            notes = (notes + "; " if notes else "") + "salary_growth_potential compares company average salary with job minimum salary"
        notes_values.append(notes)

    df["total_score"] = total_scores
    df["score_status"] = score_statuses
    df["missing_fields"] = missing_fields_values
    df["data_quality_notes"] = notes_values
    record_event(ctx, "INFO", "score_calculated", "scoring", "スコア計算が完了しました")
    return df


def build_ranking(df: pd.DataFrame) -> pd.DataFrame:
    risky_name_statuses = {"name_mismatch", "estimated_number_name_mismatch"}
    name_check_ok = ~df.get("join_name_check_status", pd.Series("", index=df.index)).isin(risky_name_statuses)
    return df[df["join_status"].eq("matched") & df["total_score"].notna() & name_check_ok].sort_values("total_score", ascending=False).reset_index(drop=True)


def dataframe_to_jsonl(df: pd.DataFrame, path: Path) -> None:
    records = df.where(pd.notna(df), None).to_dict(orient="records")
    write_jsonl(path, records)


def auto_adjust_excel(path: Path, sheet_names: list[str]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    red = PatternFill("solid", fgColor="F4CCCC")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        headers = [cell.value for cell in ws[1]]
        join_idx = headers.index("JOIN状態") + 1 if "JOIN状態" in headers else None
        score_idx = headers.index("スコア状態") + 1 if "スコア状態" in headers else None
        for row in ws.iter_rows(min_row=2):
            if join_idx and row[join_idx - 1].value not in (None, "matched"):
                for cell in row:
                    cell.fill = red
            elif score_idx and row[score_idx - 1].value not in (None, "ok"):
                for cell in row:
                    cell.fill = yellow
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_length + 2, 10), 45)
    wb.save(path)


def export_to_excel(
    df_scored: pd.DataFrame,
    df_ranking: pd.DataFrame,
    errors: list[dict[str, Any]],
    settings: dict[str, Any],
    output_path: Path,
) -> None:
    ranking_columns = [
        ("企業名", "company_name"),
        ("法人番号", "corporate_number"),
        ("法人番号状態", "corporate_number_status"),
        ("法人番号補完状態", "corporate_number_resolution_status"),
        ("法人番号補完信頼度", "corporate_number_resolution_confidence"),
        ("法人番号補完候補", "corporate_number_candidates"),
        ("法人番号補完メモ", "corporate_number_resolution_notes"),
        ("求人番号", "job_id"),
        ("総合スコア", "total_score"),
        ("研究開発スコア", "rd_score_norm"),
        ("昇給ポテンシャルスコア", "growth_score_norm"),
        ("リアル時給スコア", "wage_score_norm"),
        ("定着スコア", "tenure_score_norm"),
        ("拠点マッチスコア", "location_score_norm"),
        ("基本給下限", "basic_salary_min"),
        ("平均年間給与", "average_annual_salary"),
        ("推定リアル時給", "real_hourly_wage"),
        ("年間休日数", "annual_holidays"),
        ("月平均残業時間", "overtime_hours_avg"),
        ("研究開発費", "rd_expenses"),
        ("売上高", "net_sales"),
        ("平均勤続年数", "average_length_of_service"),
        ("就業場所", "work_location"),
        ("EDINET提出者名", "filer_name"),
        ("EDINET状態", "edinet_status"),
        ("XBRL解析状態", "xbrl_parse_status"),
        ("XBRL抽出項目", "xbrl_extracted_fields"),
        ("XBRL不足項目", "xbrl_missing_fields"),
        ("XBRL診断メモ", "xbrl_diagnostic_notes"),
        ("JOIN状態", "join_status"),
        ("JOIN失敗理由", "join_failure_reason"),
        ("JOIN診断メモ", "join_diagnostic_notes"),
        ("JOIN企業名確認", "join_name_check_status"),
        ("JOIN企業名確認メモ", "join_name_check_notes"),
        ("スコア状態", "score_status"),
        ("欠損項目", "missing_fields"),
        ("データ品質メモ", "data_quality_notes"),
    ]

    def display_frame(frame: pd.DataFrame) -> pd.DataFrame:
        out = pd.DataFrame()
        for label, column in ranking_columns:
            out[label] = frame[column] if column in frame.columns else ""
        return out

    settings_df = pd.DataFrame([{"key": key, "value": json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value} for key, value in settings.items()])
    errors_df = pd.DataFrame(errors)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        display_frame(df_ranking).to_excel(writer, sheet_name="ranking", index=False)
        display_frame(df_scored).to_excel(writer, sheet_name="all_records", index=False)
        errors_df.to_excel(writer, sheet_name="errors", index=False)
        settings_df.to_excel(writer, sheet_name="settings", index=False)
    auto_adjust_excel(output_path, ["ranking", "all_records", "errors", "settings"])


def export_outputs(
    ctx: RunContext,
    df_scored: pd.DataFrame,
    df_ranking: pd.DataFrame,
    config: dict[str, Any],
    started_at: str,
    mode: str,
    config_file: str,
    run_options: dict[str, Any] | None = None,
) -> dict[str, Any]:
    today = datetime.now().strftime("%Y%m%d")
    excel_path = ctx.output_dir / f"company_scoring_result_{today}.xlsx"
    ranking_csv = ctx.output_dir / "ranking.csv"
    all_csv = ctx.output_dir / "all_records_clean.csv"
    all_jsonl = ctx.output_dir / "all_records_clean.jsonl"
    score_components_csv = ctx.output_dir / "score_components.csv"
    errors_jsonl = ctx.output_dir / "errors.jsonl"
    manifest_path = ctx.output_dir / "run_manifest.json"

    score_columns = [
        "job_id",
        "corporate_number",
        "company_name",
        "rd_ratio",
        "salary_growth_potential",
        "real_hourly_wage",
        "tenure_score",
        "location_match_score",
        "rd_score_norm",
        "growth_score_norm",
        "wage_score_norm",
        "tenure_score_norm",
        "location_score_norm",
        "total_score",
        "score_status",
        "missing_fields",
        "join_status",
        "join_failure_reason",
        "join_name_check_status",
        "edinet_status",
        "xbrl_parse_status",
        "xbrl_extracted_fields",
        "xbrl_missing_fields",
    ]
    settings = {
        "app_name": APP_NAME,
        "app_version": APP_VERSION,
        "app_build_date": APP_BUILD_DATE,
        "run_id": ctx.run_id,
        "started_at": started_at,
        "finished_at": datetime.now().isoformat(timespec="seconds"),
        "mode": mode,
        "config_file": config_file,
        "score_weights": config["score_weights"],
        "target_locations": config["target_locations"],
        "min_valid_score_count": config["min_valid_score_count"],
        "run_options": run_options or {},
    }

    try:
        export_to_excel(df_scored, df_ranking, ctx.errors, settings, excel_path)
    except (PermissionError, zipfile.BadZipFile) as exc:
        excel_path = ctx.output_dir / f"company_scoring_result_{today}_{ctx.run_id}.xlsx"
        record_event(
            ctx,
            "WARNING",
            "excel_file_retry_with_new_name",
            "export",
            f"既存Excelを安全に更新できないため別名で保存します: {excel_path}",
            exception_type=type(exc).__name__,
            traceback=traceback.format_exc(),
        )
        export_to_excel(df_scored, df_ranking, ctx.errors, settings, excel_path)
    df_ranking.to_csv(ranking_csv, index=False, encoding="utf-8-sig")
    df_scored.to_csv(all_csv, index=False, encoding="utf-8-sig")
    dataframe_to_jsonl(df_scored, all_jsonl)
    df_scored[[column for column in score_columns if column in df_scored.columns]].to_csv(score_components_csv, index=False, encoding="utf-8-sig")
    write_jsonl(errors_jsonl, ctx.errors)

    output_files = {
        "excel": str(excel_path),
        "ranking_csv": str(ranking_csv),
        "all_records_csv": str(all_csv),
        "all_records_jsonl": str(all_jsonl),
        "score_components_csv": str(score_components_csv),
        "errors_jsonl": str(errors_jsonl),
        "log": str(ctx.log_path),
        "events_jsonl": str(ctx.events_path),
    }
    manifest = {
        **settings,
        "input_count": int(len(df_scored)),
        "ranking_count": int(len(df_ranking)),
        "join_matched_count": int(df_scored["join_status"].eq("matched").sum()),
        "join_failed_count": int((~df_scored["join_status"].eq("matched")).sum()),
        "join_name_mismatch_count": int(df_scored.get("join_name_check_status", pd.Series("", index=df_scored.index)).isin({"name_mismatch", "estimated_number_name_mismatch"}).sum()),
        "score_ok_count": int(df_scored["score_status"].eq("ok").sum()),
        "score_partial_count": int(df_scored["score_status"].eq("partial").sum()),
        "score_insufficient_count": int(df_scored["score_status"].eq("insufficient_data").sum()),
        "error_count": int(sum(1 for item in ctx.errors if item.get("level") == "ERROR")),
        "warning_count": int(sum(1 for item in ctx.errors if item.get("level") == "WARNING")),
        "output_files": output_files,
    }
    with manifest_path.open("w", encoding="utf-8-sig") as handle:
        json.dump(manifest, handle, ensure_ascii=False, indent=2)
    output_files["manifest"] = str(manifest_path)
    record_event(ctx, "INFO", "export_completed", "export", "出力が完了しました")
    return output_files


def load_input_data(ctx: RunContext) -> tuple[pd.DataFrame, pd.DataFrame]:
    if USE_EXISTING_CLEAN_DATA:
        raise NotImplementedError("USE_EXISTING_CLEAN_DATA は次の試作品で実装します")
    if USE_MOCK_DATA:
        record_event(ctx, "INFO", "input_loaded", "load", "モックデータを読み込みました")
        return load_mock_jobs(), load_mock_edinet()
    if USE_JOB_FILE_INPUT:
        raise NotImplementedError("CSV/Excel入力は次の試作品で実装します")
    if USE_HELLOWORK_API:
        return fetch_jobs_from_hellowork(ctx), pd.DataFrame()
    raise ValueError("データソース設定がすべてFalseです")


def main() -> None:
    args = parse_args()
    started_at = datetime.now().isoformat(timespec="seconds")
    output_dir = Path(args.output_dir)
    ctx = RunContext(make_run_id(), output_dir)
    setup_logging(ctx)
    try:
        config = load_config(ctx, args.config)
        if args.rescore:
            existing_path = resolve_existing_clean_data_path(args.input_clean_data, output_dir)
            df_merged = load_existing_clean_data(existing_path, ctx)
            record_event(ctx, "INFO", "rescore_started", "scoring", "API通信なしの再スコアリングを開始しました")
        elif args.job_file:
            job_path = Path(args.job_file)
            df_jobs_raw = load_jobs_from_file(job_path, ctx)
            if args.corporate_master:
                corporate_master = load_corporate_master(Path(args.corporate_master), ctx)
                df_jobs_raw = apply_corporate_number_master(df_jobs_raw, corporate_master, ctx)
            if args.use_edinet_api:
                target_dates = make_recent_date_strings(args.edinet_lookback_days)
                df_edinet_raw = fetch_edinet_document_list(target_dates, ctx)
                if df_edinet_raw.empty:
                    record_event(ctx, "WARNING", "edinet_fallback_mock_used", "load", "EDINET取得結果が空のためモックデータを使用します")
                    df_edinet_raw = load_mock_edinet()
                elif args.edinet_dry_run:
                    record_event(ctx, "INFO", "edinet_dry_run", "edinet", "EDINETドライランのためXBRL ZIP取得をスキップしました")
                    export_edinet_snapshot(df_edinet_raw, ctx)
                else:
                    target_numbers = {
                        normalize_corporate_number(value)
                        for value in df_jobs_raw.get("corporate_number", pd.Series(dtype=str)).tolist()
                        if normalize_corporate_number(value)
                    }
                    df_edinet_raw = enrich_edinet_records_with_xbrl(df_edinet_raw, target_numbers, ctx, args.edinet_xbrl_limit)
                    export_edinet_snapshot(df_edinet_raw, ctx)
            else:
                df_edinet_raw = load_mock_edinet()
                record_event(ctx, "WARNING", "edinet_mock_used", "load", "求人ファイル入力の検証用にEDINETはモックデータを使用します")
            df_jobs = clean_jobs(df_jobs_raw, ctx)
            df_edinet = clean_edinet(df_edinet_raw, ctx)
            df_merged = merge_jobs_and_edinet(df_jobs, df_edinet)
        else:
            df_jobs_raw, df_edinet_raw = load_input_data(ctx)
            df_jobs = clean_jobs(df_jobs_raw, ctx)
            df_edinet = clean_edinet(df_edinet_raw, ctx)
            df_merged = merge_jobs_and_edinet(df_jobs, df_edinet)
        df_scored = calculate_score_columns(df_merged, config, ctx)
        df_ranking = build_ranking(df_scored)
        mode = "rescore" if args.rescore else ("file_input" if args.job_file else ("mock" if USE_MOCK_DATA else "other"))
        run_options = {
            "job_file": args.job_file,
            "corporate_master": args.corporate_master,
            "use_edinet_api": args.use_edinet_api,
            "edinet_dry_run": args.edinet_dry_run,
            "edinet_lookback_days": args.edinet_lookback_days,
            "edinet_xbrl_limit": args.edinet_xbrl_limit,
            "rescore": args.rescore,
            "input_clean_data": args.input_clean_data,
        }
        outputs = export_outputs(ctx, df_scored, df_ranking, config, started_at, mode, args.config, run_options)
        logging.info("完了しました。主な出力: %s", json.dumps(outputs, ensure_ascii=False))
    except Exception as exc:
        record_event(
            ctx,
            "ERROR",
            "fatal_error",
            "main",
            "処理を停止しました",
            exception_type=type(exc).__name__,
            traceback=traceback.format_exc(),
        )
        raise


if __name__ == "__main__":
    main()
